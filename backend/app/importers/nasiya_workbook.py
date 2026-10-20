"""Read a daily "Telefon nasiya nazorat" .xlsx into the pure parser (app.domain.nasiya_sheet). See docs/SPEC.md §18."""

import io
import zipfile
from datetime import date

from openpyxl import load_workbook

from app.core.errors import ValidationFailed
from app.domain import nasiya_sheet as ns

MAX_BYTES = 20 * 1024 * 1024  # Telegram bots cannot download bigger files anyway
MAX_ROWS = 2000
MAX_COLS = 60


def looks_like_xlsx(file_name: str | None) -> bool:
    return bool(file_name) and file_name.lower().endswith((".xlsx", ".xlsm"))


def _spread_merged_headers(ws, grid: list[list]) -> None:
    """A header merged over several columns names every one of them. Only header rows: a merged data cell
    must not turn into several IMEIs."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row > ns.HEADER_SCAN_ROWS or rng.min_row != rng.max_row:
            continue
        r = rng.min_row - 1
        if r >= len(grid):
            continue
        value = grid[r][rng.min_col - 1] if rng.min_col - 1 < len(grid[r]) else None
        if not isinstance(value, str) or ns.column_kind(value) is None:
            continue
        for c in range(rng.min_col, min(rng.max_col, MAX_COLS)):
            if c < len(grid[r]) and grid[r][c] is None:
                grid[r][c] = value


def parse_workbook(content: bytes, file_name: str = "", fallback_date: date | None = None) -> ns.WorkbookResult:
    if len(content) > MAX_BYTES:
        raise ValidationFailed("Fayl juda katta (20 MB dan ortiq)")
    try:
        # Not read_only: merged header cells ("Насияга берилди" over two columns) are needed.
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ValidationFailed("Excel faylni o'qib bo'lmadi") from exc
    named_date = ns.date_from_filename(file_name)
    default_year = (named_date or fallback_date or date.today()).year
    sheets = []
    try:
        for ws in wb.worksheets:
            grid = [list(row[:MAX_COLS]) for row in ws.iter_rows(max_row=MAX_ROWS, values_only=True)]
            _spread_merged_headers(ws, grid)
            sheets.append(ns.parse_sheet(ws.title, grid, default_year))
    finally:
        wb.close()
    sheets = [s for s in sheets if s.has_data or s.summary]
    if not sheets:
        raise ValidationFailed("Faylda nasiya nazorati varaqlari topilmadi (Nasiya/Qoldiq/Sotuv ustunlari yo'q)")
    return ns.WorkbookResult(report_date=ns.report_date(sheets, named_date or fallback_date), sheets=sheets,
                             issues=ns.check_workbook(sheets))
