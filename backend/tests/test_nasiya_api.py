"""Daily nasiya control (SPEC §18): upload / Telegram intake, review, confirmation and monitoring."""

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from sqlalchemy import select

from app.models import NasiyaIntake, NasiyaReport, TelegramGroupMessage
from app.services import nasiya_telegram
from tests.test_customer_accounts import _customer
from tests.test_nasiya_sheet import imei


def workbook(day: str, sheets: dict[str, list[list]]) -> bytes:
    """sheets: name → columns (each a list of cells top-down); row 1 gets the date."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, columns in sheets.items():
        ws = wb.create_sheet(name)
        ws.cell(1, 1, day)
        for c, column in enumerate(columns, start=1):
            for r, value in enumerate(column, start=2):
                if value is not None:
                    ws.cell(r, c, value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DAY1 = {
    "Ali aka": [
        ["Nasiya", "17 max 256", int(imei(1)), int(imei(2)), "1500*2", "13 black 128", int(imei(3)), 460],
        ["Qoldiq", "16 black 128", int(imei(4)), int(imei(5)), 2],
    ],
    "Vali aka": [["Qoldiq", "15 blue 128", int(imei(6)), int(imei(7)), 2]],
}
# Day 2: Ali paid imei 1 (sotuv), imei 2 vanished from his debt, imei 4 sold on credit, imei 8 is new debt;
# Vali's imei 7 moved to Ali's qoldiq.
DAY2 = {
    "Ali aka": [
        ["Sotuv", "17 max 256", int(imei(1)), 1500],
        ["Nasiya", "13 black 128", int(imei(3)), 460, "16 black 128", int(imei(4)), 790, "18 pro 256", int(imei(8)), 1600],
        ["Qoldiq", "16 black 128", int(imei(5)), 1, "15 blue 128", int(imei(7)), 1],
    ],
    "Vali aka": [["Qoldiq", "15 blue 128", int(imei(6)), 1]],
}


async def _upload(client, h, content: bytes, name: str, expect=201) -> dict:
    r = await client.post("/api/nasiya/reports/upload", files={"file": (name, content)}, headers=h)
    assert r.status_code == expect, r.text
    return r.json()


async def _detail(client, h, report_id: int) -> dict:
    r = await client.get(f"/api/nasiya/reports/{report_id}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


async def _link_all(client, h, detail: dict, customers: dict[str, int]) -> dict:
    for sheet in detail["sheets"]:
        if sheet["customer"] is None:
            r = await client.put(f"/api/nasiya/reports/{detail['id']}/sheets/{sheet['id']}",
                                 json={"customer_id": customers[sheet["name"]]}, headers=h)
            assert r.status_code == 200, r.text
            detail = r.json()
    return detail


async def _confirm(client, h, report_id: int, expect=200, **body) -> dict:
    r = await client.post(f"/api/nasiya/reports/{report_id}/confirm", json=body, headers=h)
    assert r.status_code == expect, r.text
    return r.json()


async def _setup_two_days(client, h):
    ali = await _customer(client, h, "Ali aka")
    vali = await _customer(client, h, "Vali aka")
    customers = {"Ali aka": ali["id"], "Vali aka": vali["id"]}
    up1 = await _upload(client, h, workbook("14,10,2026", DAY1), "nazorat_14.10.2026.xlsx")
    d1 = await _link_all(client, h, await _detail(client, h, up1["report_id"]), customers)
    await _confirm(client, h, d1["id"])
    up2 = await _upload(client, h, workbook("15,10,2026", DAY2), "nazorat_15.10.2026.xlsx")
    d2 = await _detail(client, h, up2["report_id"])
    return customers, d1, d2


async def test_upload_creates_a_pending_report_with_items_and_suggestions(client, operator_h):
    ali = await _customer(client, operator_h, "Ali aka")
    content = workbook("14,10,2026", DAY1)  # one save: openpyxl stamps the time, a second save can differ
    up = await _upload(client, operator_h, content, "nazorat_14.10.2026.xlsx")
    assert up["report_date"] == "2026-10-14" and up["duplicate"] is False and up["error"] is None
    d = await _detail(client, operator_h, up["report_id"])
    assert (d["status"], d["source"], d["item_count"], d["error_count"], d["unlinked_sheets"]) == (
        "pending", "upload", 7, 0, 2)
    ali_sheet = d["sheets"][0]
    assert ali_sheet["name"] == "Ali aka" and ali_sheet["suggested_customer"]["id"] == ali["id"]
    assert {t["kind"]: (t["qty"], t["amount"]) for t in ali_sheet["totals"]} == {
        "nasiya": (3, 3460), "qoldiq": (2, 0)}
    assert [(i["model"], i["imei"], i["price_usd"]) for i in ali_sheet["items"]][:3] == [
        ("17 max 256", imei(1), 1500), ("17 max 256", imei(2), 1500), ("13 black 128", imei(3), 460)]
    assert d["diff"] is None  # nothing confirmed before it

    again = await _upload(client, operator_h, content, "copy.xlsx")
    assert again["duplicate"] is True and again["report_id"] == up["report_id"]


async def test_upload_refuses_a_non_nasiya_file(client, operator_h):
    up = await _upload(client, operator_h, workbook("14,10,2026", {"Hisob": [["Ism", "Summa"]]}), "x.xlsx")
    assert up["report_id"] is None and "topilmadi" in up["error"]
    up = await _upload(client, operator_h, b"not an excel file", "x.xlsx")
    assert up["error"] == "Excel faylni o'qib bo'lmadi"


async def test_confirmation_needs_every_sheet_linked_and_no_errors(client, operator_h, admin_h):
    ali = await _customer(client, operator_h, "Ali aka")
    bad = {"Ali aka": [["Nasiya", "17 max 256", int(imei(1)), "1500*2"]]}  # 2 written, 1 IMEI
    up = await _upload(client, operator_h, workbook("14,10,2026", bad), "n.xlsx")
    r = await client.post(f"/api/nasiya/reports/{up['report_id']}/confirm", json={}, headers=operator_h)
    assert r.status_code == 422 and "bog'lanmagan" in r.json()["detail"]
    d = await _link_all(client, operator_h, await _detail(client, operator_h, up["report_id"]), {"Ali aka": ali["id"]})
    assert d["error_count"] == 1
    r = await client.post(f"/api/nasiya/reports/{d['id']}/confirm", json={}, headers=operator_h)
    assert r.status_code == 422 and "1 ta xato" in r.json()["detail"]
    r = await client.post(f"/api/nasiya/reports/{d['id']}/confirm", json={"force": True}, headers=operator_h)
    assert r.status_code == 403
    done = await _confirm(client, admin_h, d["id"], force=True, note="IMEI keyin tuzatiladi")
    assert done["status"] == "confirmed" and done["review_note"] == "IMEI keyin tuzatiladi"
    assert done["reviewed_by"] == "t_admin"


async def test_a_sheet_link_is_remembered_for_the_next_file(client, operator_h):
    customers, d1, d2 = await _setup_two_days(client, operator_h)
    assert [s["customer"]["id"] for s in d2["sheets"]] == [customers["Ali aka"], customers["Vali aka"]]
    assert d2["unlinked_sheets"] == 0
    # one customer cannot take two sheets of the same file
    r = await client.put(f"/api/nasiya/reports/{d2['id']}/sheets/{d2['sheets'][1]['id']}",
                         json={"customer_id": customers["Ali aka"]}, headers=operator_h)
    assert r.status_code == 409


async def test_day_to_day_events_by_imei(client, operator_h):
    customers, d1, d2 = await _setup_two_days(client, operator_h)
    diff = d2["diff"]
    assert diff["previous_date"] == "2026-10-14"
    counts = {c["type"]: (c["qty"], c["amount"], c["alert"]) for c in diff["counts"]}
    assert counts == {
        "vanished_debt": (1, 1500, True),   # imei 2: left Ali's debt without being paid
        "moved": (1, 0, True),              # imei 7: Vali's qoldiq → Ali's qoldiq
        "paid": (1, 1500, False),           # imei 1: nasiya → sotuv
        "qoldiq_to_debt": (1, 790, False),  # imei 4: sold on credit
        "new_debt": (1, 1600, False),       # imei 8
    }
    moved = next(e for e in diff["events"] if e["type"] == "moved")
    assert (moved["imei"], moved["customer"]["name"], moved["other_customer"]["name"]) == (
        imei(7), "Ali aka", "Vali aka")


async def test_overview_after_two_confirmed_days(client, operator_h, viewer_h):
    customers, d1, d2 = await _setup_two_days(client, operator_h)
    pending = (await client.get("/api/nasiya/overview", headers=viewer_h)).json()
    assert pending["report"]["id"] == d1["id"] and pending["pending_reports"] == 1
    await _confirm(client, operator_h, d2["id"])

    o = (await client.get("/api/nasiya/overview", headers=viewer_h)).json()
    assert o["report"]["report_date"] == "2026-10-15" and o["previous_date"] == "2026-10-14"
    assert o["pending_reports"] == 0 and o["history_start"] == "2026-10-14"
    t = o["totals"]
    assert (t["debt_usd"], t["debt_qty"], t["qoldiq_qty"], t["sotuv_usd"], t["sotuv_qty"]) == (2850, 3, 3, 1500, 1)
    assert o["previous_totals"]["debt_usd"] == 3460
    assert {a["type"] for a in o["alerts"]} == {"vanished_debt", "moved"}
    ali = next(r for r in o["customers"] if r["customer"]["name"] == "Ali aka")
    assert (ali["totals"]["debt_usd"], ali["new_debt_usd"], ali["paid_usd"], ali["alert_count"]) == (
        2850, 2390, 1500, 2)
    assert ali["panel_debt_usd"] == 0  # nothing in the panel ledger: the difference is visible
    # imei 3 has been in Ali's debt for both confirmed days; ages start at the first confirmed day (lower bound)
    assert ali["max_age_days"] == 1
    assert [b["qty"] for b in o["aging"]] == [3, 0, 0, 0, 0]


async def test_customer_detail_lists_every_phone_with_model_and_age(client, operator_h):
    customers, d1, d2 = await _setup_two_days(client, operator_h)
    await _confirm(client, operator_h, d2["id"])
    r = await client.get(f"/api/nasiya/customers/{customers['Ali aka']}", headers=operator_h)
    assert r.status_code == 200, r.text
    c = r.json()
    rows = [(i["kind"], i["model"], i["imei"], i["price_usd"], i["since"], i["since_exact"], i["age_days"])
            for i in c["items"]]
    assert rows == [
        ("nasiya", "13 black 128", imei(3), 460, "2026-10-14", False, 1),
        ("nasiya", "16 black 128", imei(4), 790, "2026-10-15", True, 0),
        ("nasiya", "18 pro 256", imei(8), 1600, "2026-10-15", True, 0),
        ("qoldiq", "16 black 128", imei(5), None, "2026-10-14", False, 1),
        ("qoldiq", "15 blue 128", imei(7), None, "2026-10-15", True, 0),
        ("sotuv", "17 max 256", imei(1), 1500, None, False, None),
    ]
    assert [(m["kind"], m["model"], m["qty"], m["amount"]) for m in c["models"]] == [
        ("nasiya", "13 black 128", 1, 460), ("nasiya", "16 black 128", 1, 790), ("nasiya", "18 pro 256", 1, 1600),
        ("qoldiq", "15 blue 128", 1, 0), ("qoldiq", "16 black 128", 1, 0), ("sotuv", "17 max 256", 1, 1500)]
    assert [p["debt_usd"] for p in c["series"]] == [3460, 2850]

    trail = (await client.get(f"/api/nasiya/imei/{imei(4)}", headers=operator_h)).json()
    assert [(t["date"], t["kind"], t["customer"]["name"]) for t in trail] == [
        ("2026-10-15", "nasiya", "Ali aka"), ("2026-10-14", "qoldiq", "Ali aka")]


async def test_a_newer_file_for_the_same_day_supersedes_the_confirmed_one(client, operator_h):
    customers, d1, _ = await _setup_two_days(client, operator_h)
    fixed = {**DAY1, "Vali aka": [["Qoldiq", "15 blue 128", int(imei(6)), 1]]}
    up = await _upload(client, operator_h, workbook("14,10,2026", fixed), "nazorat_14.10.2026 (2).xlsx")
    await _confirm(client, operator_h, up["report_id"])
    assert (await _detail(client, operator_h, d1["id"]))["status"] == "superseded"
    listed = (await client.get("/api/nasiya/reports?status=confirmed", headers=operator_h)).json()
    assert [r["id"] for r in listed] == [up["report_id"]]


async def test_reject_and_delete(client, operator_h, admin_h):
    up = await _upload(client, operator_h, workbook("14,10,2026", DAY1), "n.xlsx")
    r = await client.post(f"/api/nasiya/reports/{up['report_id']}/reject", json={"note": "eski fayl"},
                          headers=operator_h)
    assert r.status_code == 200 and r.json()["status"] == "rejected"
    assert (await client.delete(f"/api/nasiya/reports/{up['report_id']}", headers=operator_h)).status_code == 403
    assert (await client.delete(f"/api/nasiya/reports/{up['report_id']}", headers=admin_h)).status_code == 204
    assert (await client.get(f"/api/nasiya/reports/{up['report_id']}", headers=admin_h)).status_code == 404


async def test_roles(client, viewer_h):
    assert (await client.get("/api/nasiya/overview", headers=viewer_h)).status_code == 200
    r = await client.post("/api/nasiya/reports/upload", files={"file": ("n.xlsx", b"x")}, headers=viewer_h)
    assert r.status_code == 403
    assert (await client.get("/api/nasiya/overview")).status_code == 401


async def _group_message(session, message_id: int, file_name: str, chat_id: int = -1003839931967):
    row = TelegramGroupMessage(chat_id=chat_id, chat_title="Telefon nasiya nazirat", message_id=message_id,
                               from_id=5005, from_name="Ali", sent_at=datetime(2026, 10, 15, 18, tzinfo=timezone.utc),
                               content_type="document", text=None, file_id=f"FILE{message_id}", file_name=file_name,
                               raw={"document": {"file_size": 40000}})
    session.add(row)
    await session.flush()
    return row


async def test_telegram_documents_are_ingested_and_rescanned_once(session, monkeypatch):
    day2 = workbook("15,10,2026", DAY2)  # one save: the same bytes twice
    files = {"FILE1": day2, "FILE2": b"not excel", "FILE3": day2}
    downloads = []

    async def download(file_id):
        downloads.append(file_id)
        return files[file_id]

    row = await _group_message(session, 1, "nazorat_15.10.2026.xlsx")
    assert await nasiya_telegram.ingest_message(session, row, download) == "created"
    report = (await session.execute(select(NasiyaReport).where(NasiyaReport.telegram_message_id == row.id))).scalar_one()
    assert (report.status, report.source, report.report_date.isoformat(), report.sent_by, report.chat_title) == (
        "pending", "telegram", "2026-10-15", "Ali", "Telefon nasiya nazirat")
    assert await nasiya_telegram.ingest_message(session, row, download) is None  # handled once
    await _group_message(session, 2, "boshqa.xlsx")
    await _group_message(session, 3, "copy.xlsx")  # the same workbook sent again
    await _group_message(session, 4, "photo.jpg")
    result = await nasiya_telegram.rescan(session, download)
    assert (result.checked, result.created, result.skipped, result.failed) == (2, 0, 2, 0)  # 1 handled, 4 not .xlsx
    outcomes = dict((await session.execute(select(NasiyaIntake.telegram_message_id, NasiyaIntake.outcome))).all())
    assert sorted(outcomes.values()) == ["created", "duplicate", "invalid"]
    invalid = (await session.execute(select(NasiyaReport).where(NasiyaReport.status == "invalid"))).scalar_one()
    assert invalid.issues[0]["code"] == "unreadable"
    assert downloads == ["FILE1", "FILE2", "FILE3"]
    again = await nasiya_telegram.rescan(session, download)
    assert again.checked == 0 and downloads == ["FILE1", "FILE2", "FILE3"]

    monkeypatch.setattr(nasiya_telegram, "nasiya_chat_ids", lambda: {-100})
    assert not nasiya_telegram.is_candidate(-1003839931967, "n.xlsx")
    assert nasiya_telegram.is_candidate(-100, "n.XLSX") and not nasiya_telegram.is_candidate(-100, "n.pdf")


async def _process(client, h, since: str | None = None) -> dict:
    r = await client.post("/api/nasiya/reports/process", json={"since": since}, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


async def test_process_confirms_every_day_from_a_date(client, operator_h, admin_h):
    ali = await _customer(client, operator_h, "Ali aka")
    vali = await _customer(client, operator_h, "Vali aka")
    old = await _upload(client, operator_h, workbook("13,10,2026", DAY1), "nazorat_13.10.2026.xlsx")
    d1 = await _upload(client, operator_h, workbook("14,10,2026", DAY1), "nazorat_14.10.2026.xlsx")
    first_try = {**DAY2, "Vali aka": [["Qoldiq", "15 blue 128", int(imei(6)), 1, "x"]]}
    early = await _upload(client, operator_h, workbook("15,10,2026", first_try), "nazorat_15.10.2026.xlsx")
    d2 = await _upload(client, operator_h, workbook("15,10,2026", DAY2), "nazorat_15.10.2026 (2).xlsx")
    bad = {**DAY2, "Vali aka": [["Qoldiq", "15 blue 128", int(imei(6)), 1], ["Nasiya", "14 128", int(imei(9)), "900*2"]]}
    d3 = await _upload(client, operator_h, workbook("16,10,2026", bad), "nazorat_16.10.2026.xlsx")

    out = await _process(client, operator_h, "2026-10-14")
    assert (out["confirmed"], out["superseded"]) == (2, 1)
    assert sorted(out["linked"]) == ["Ali aka → Ali aka", "Vali aka → Vali aka"]
    assert [(w["report_id"], w["reason"]) for w in out["waiting"]] == [
        (d3["report_id"], "1 ta xato: administrator tasdiqlashi kerak")]
    status = {r["id"]: r["status"] for r in (await client.get("/api/nasiya/reports", headers=operator_h)).json()}
    assert status == {old["report_id"]: "pending", d1["report_id"]: "confirmed", early["report_id"]: "superseded",
                      d2["report_id"]: "confirmed", d3["report_id"]: "pending"}

    months = (await client.get("/api/nasiya/months", headers=operator_h)).json()
    assert len(months) == 1
    m = months[0]
    assert (m["month"], m["days"], m["first_date"], m["last_date"], m["partial"]) == (
        "2026-10-01", 2, "2026-10-14", "2026-10-15", True)
    # 15.10: imei 1 paid (Sotuv); imei 4 (qoldiq → nasiya) and imei 8 are new debt
    assert (m["sotuv_qty"], m["sotuv_usd"], m["new_debt_qty"], m["new_debt_usd"]) == (1, 1500, 2, 2390)
    assert (m["debt_qty_end"], m["debt_usd_end"], m["qoldiq_qty_end"]) == (3, 2850, 3)
    vali_months = (await client.get(f"/api/nasiya/months?customer_id={vali['id']}", headers=operator_h)).json()
    assert (vali_months[0]["debt_qty_end"], vali_months[0]["qoldiq_qty_end"]) == (0, 1)

    out = await _process(client, admin_h, "2026-10-14")  # an administrator takes the file with errors too
    assert (out["confirmed"], out["waiting"]) == (1, [])
    assert (await _detail(client, admin_h, d3["report_id"]))["review_note"] == "Hammasini hisoblash: 1 ta xato bilan"
    assert ali["id"] != vali["id"]


async def test_a_group_file_of_known_customers_is_counted_on_arrival(session):
    from app.models import Customer

    session.add_all([Customer(name="Ali aka", kind="dealer"), Customer(name="Vali aka", kind="dealer")])
    files = {"FILE1": workbook("14,10,2026", DAY1), "FILE2": workbook("15,10,2026", DAY2),
             "FILE3": workbook("16,10,2026", {"Ali aka": [["Nasiya", "17 max 256", int(imei(1)), "1500*2"]]}),
             "FILE4": workbook("17,10,2026", {"Yangi odam": [["Nasiya", "17 max 256", int(imei(1)), 1500]]})}

    async def download(file_id):
        return files[file_id]

    rows = [await _group_message(session, n, f"nazorat_{14 + n - 1}.10.2026.xlsx") for n in (1, 2, 3, 4)]
    rows[0].sent_at = datetime(2026, 10, 13, 18, tzinfo=timezone.utc)  # before the requested day: skipped
    await session.flush()
    result = await nasiya_telegram.rescan(session, download, since=datetime(2026, 10, 14).date(), limit=1)
    assert (result.checked, result.created, result.confirmed, result.remaining) == (1, 1, 1, 2)
    result = await nasiya_telegram.rescan(session, download, since=datetime(2026, 10, 14).date())
    assert (result.checked, result.created, result.confirmed, result.remaining) == (2, 2, 1, 0)
    reports = {r.report_date.isoformat(): r for r in (await session.execute(select(NasiyaReport))).scalars()}
    assert {d: r.status for d, r in reports.items()} == {
        "2026-10-15": "confirmed", "2026-10-16": "confirmed", "2026-10-17": "pending"}  # no customer matches
    assert (reports["2026-10-15"].review_note, reports["2026-10-16"].review_note) == (
        "Avtomatik: xatosiz", "Avtomatik: 1 ta xato bilan")


async def test_daily_cron_needs_the_secret_and_counts_pending_days(client, operator_h, monkeypatch):
    from app.core.config import get_settings

    await _customer(client, operator_h, "Ali aka")
    await _customer(client, operator_h, "Vali aka")
    up = await _upload(client, operator_h, workbook("14,10,2026", DAY1), "nazorat_14.10.2026.xlsx")
    assert (await client.get("/api/nasiya/cron")).status_code == 403  # no CRON_SECRET: off
    monkeypatch.setattr(get_settings(), "cron_secret", "s" * 32)
    monkeypatch.setattr(get_settings(), "telegram_bot_token", "")
    assert (await client.get("/api/nasiya/cron", headers={"Authorization": "Bearer wrong"})).status_code == 403
    r = await client.get("/api/nasiya/cron", headers={"Authorization": "Bearer " + "s" * 32})
    assert r.status_code == 200, r.text
    assert r.json()["processed"]["confirmed"] == 1
    d = await _detail(client, operator_h, up["report_id"])
    assert (d["status"], d["review_note"]) == ("confirmed", "Kunlik avtomatik hisoblash")


class _FakeBot:
    """Answers the read calls bot_status makes."""

    def __init__(self, *, privacy_off: bool, member_status: str, webhook_url: str = "https://x/api/telegram/webhook"):
        from types import SimpleNamespace as NS

        self._me = NS(id=42, username="teztel_bot", can_read_all_group_messages=privacy_off)
        self._webhook = NS(url=webhook_url, last_error_message=None, pending_update_count=0,
                           allowed_updates=["message", "edited_message"])
        self._member = NS(status=member_status)
        self._chat = NS(title="Telefon nasiya nazirat")

    async def get_me(self):
        return self._me

    async def get_webhook_info(self):
        return self._webhook

    async def get_chat_member(self, chat_id, user_id):
        return self._member

    async def get_chat(self, chat_id):
        return self._chat


async def test_bot_status_explains_why_group_files_do_not_arrive(session, monkeypatch):
    monkeypatch.setattr(nasiya_telegram, "nasiya_chat_ids", lambda: {-1003839931967})
    status = await nasiya_telegram.bot_status(session, _FakeBot(privacy_off=False, member_status="member"))
    assert [g.title for g in status.groups] == ["Telefon nasiya nazirat"] and status.groups[0].messages == 0
    assert len(status.problems) == 1 and "administrator qiling" in status.problems[0]

    status = await nasiya_telegram.bot_status(session, _FakeBot(privacy_off=False, member_status="administrator"))
    assert status.problems == ["«Telefon nasiya nazirat»: bot ulangandan beri guruhga Excel fayl tashlanmagan."]

    await _group_message(session, 1, "nazorat_15.10.2026.xlsx")
    status = await nasiya_telegram.bot_status(session, _FakeBot(privacy_off=False, member_status="member"))
    assert status.problems == [] and status.groups[0].excel_files == 1  # files do arrive

    status = await nasiya_telegram.bot_status(session, _FakeBot(privacy_off=True, member_status="left", webhook_url=""))
    assert status.problems[0].startswith("Bot webhook o'rnatilmagan") and "guruhda emas" in status.problems[1]
